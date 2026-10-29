#!/usr/bin/env python3
"""根据草稿 JSON 生成排版整齐的 Word 文档。

草稿 JSON 结构：
{
  "title": "文章标题",
  "lead": "导语",
  "sections": [{"subtitle": "小标题", "body": "正文段落"}],
  "images": [{"caption": "图注", "source": "来源", "license": "授权说明", "path": "可选本地图片路径"}],
  "sources": ["来源链接1", "来源链接2"],
  "audit": ["审核项1", "审核项2"]   # 可选，缺省使用默认清单
}
"""
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("缺少 python-docx，请先执行: pip install python-docx")
    sys.exit(1)


DEFAULT_AUDIT = [
    "事实已多源核实",
    "无编造数据/引用/阅读量",
    "图片来源与授权明确",
    "标题与正文一致",
    "不涉及任何未公开凭证",
]


def build_doc(draft, out_path):
    doc = Document()

    title = doc.add_heading(draft.get("title", "未命名文章"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if draft.get("lead"):
        p = doc.add_paragraph()
        r = p.add_run(draft["lead"])
        r.italic = True

    for sec in draft.get("sections", []):
        if sec.get("subtitle"):
            doc.add_heading(sec["subtitle"], level=1)
        if sec.get("body"):
            doc.add_paragraph(sec["body"])

    for img in draft.get("images", []):
        caption = img.get("caption", "")
        source = img.get("source", "")
        license_info = img.get("license", "")
        doc.add_paragraph(f"[配图：{caption}]")
        note = doc.add_paragraph()
        note.add_run(f"图注：{caption}　来源：{source}　授权：{license_info}").italic = True
        path = img.get("path")
        if path and Path(path).exists():
            try:
                doc.add_picture(path, width=Inches(5.5))
            except Exception as e:
                doc.add_paragraph(f"（图片插入失败：{e}）")

    if draft.get("sources"):
        doc.add_heading("资料来源", level=1)
        for src in draft["sources"]:
            doc.add_paragraph(src, style="List Bullet")

    doc.add_heading("发布前审核清单", level=1)
    for item in draft.get("audit", DEFAULT_AUDIT):
        doc.add_paragraph(item, style="List Bullet")

    doc.save(out_path)
    return out_path


def main(argv):
    if len(argv) < 3:
        print("用法: python generate_word.py <draft.json> <out.docx>")
        return 1
    draft_path, out_path = argv[1], argv[2]
    if not Path(draft_path).exists():
        print(f"找不到草稿: {draft_path}")
        return 1
    draft = json.loads(Path(draft_path).read_text(encoding="utf-8"))
    result = build_doc(draft, out_path)
    print(f"已生成 Word 文档: {result}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
